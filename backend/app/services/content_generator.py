"""
ContentForge AI - Content Transformation Service (Phase 5)
Generates format-specific content from extracted source text using:
- Google Gemini API (gemini-2.5-flash, gemini-2.0-flash, gemini-1.5-flash)
- python-docx for downloadable .docx documents
- python-pptx for downloadable 16:9 widescreen presentations
- Image prompt optimizer connected to AI image generation
- Robust deterministic fallback when Gemini API key is missing or offline
"""

import os
import io
import re
import time
import json
import uuid
import base64
import urllib.parse
from typing import Dict, Any, List, Optional
from datetime import datetime
import requests
from dotenv import load_dotenv

import docx
from docx.shared import Inches, Pt, RGBColor

import pptx
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor

load_dotenv()

# In-memory storage for generated downloadable files
# Key: file_id (str), Value: {"bytes": bytes, "filename": str, "mime_type": str, "created_at": float}
FILE_STORE: Dict[str, Dict[str, Any]] = {}
FILE_TTL_SECONDS = 3600  # 1 hour retention


def cleanup_file_store():
    """Removes expired generated files from in-memory cache."""
    now = time.time()
    expired = [fid for fid, data in FILE_STORE.items() if now - data.get("created_at", 0) > FILE_TTL_SECONDS]
    for fid in expired:
        FILE_STORE.pop(fid, None)


def store_file(file_bytes: bytes, filename: str, mime_type: str) -> str:
    """Stores binary file in memory and returns a unique file_id."""
    cleanup_file_store()
    file_id = str(uuid.uuid4())
    FILE_STORE[file_id] = {
        "bytes": file_bytes,
        "filename": filename,
        "mime_type": mime_type,
        "created_at": time.time()
    }
    return file_id


def get_stored_file(file_id: str) -> Optional[Dict[str, Any]]:
    """Retrieves file binary and metadata by file_id."""
    cleanup_file_store()
    return FILE_STORE.get(file_id)


def clean_json_response(raw_text: str) -> str:
    """Strips markdown code fences and extraneous text from JSON response."""
    text = raw_text.strip()
    if text.startswith("```"):
        match = re.search(r'```(?:json)?\s*([\s\S]*?)\s*```', text)
        if match:
            text = match.group(1).strip()
    return text


def extract_quick_metadata(source_text: str) -> Dict[str, Any]:
    """Helper to extract clean topics and sentences for deterministic fallback generation."""
    clean = " ".join(source_text.split())
    sentences = [s.strip() for s in re.split(r'[.!?]+', clean) if len(s.strip()) > 15]
    words = re.findall(r'\b[A-Za-z]{4,}\b', clean)
    
    # Word frequencies for keyword extraction
    freq: Dict[str, int] = {}
    stopwords = {"with", "that", "this", "from", "they", "have", "were", "been", "their", "which", "about", "there", "would"}
    for w in words:
        wl = w.lower()
        if wl not in stopwords:
            freq[wl] = freq.get(wl, 0) + 1
    sorted_words = sorted(freq.items(), key=lambda x: x[1], reverse=True)
    top_keywords = [w[0].title() for w in sorted_words[:10]]
    
    main_topic = top_keywords[0] if top_keywords else "Document"
    first_sentence = sentences[0] if sentences else (source_text[:120] + "...")
    
    return {
        "sentences": sentences,
        "keywords": top_keywords,
        "main_topic": main_topic,
        "first_sentence": first_sentence
    }


# ==============================================================================
# Gemini LLM Invoker
# ==============================================================================

def call_gemini_api(prompt: str, json_mode: bool = True, temperature: float = 0.2) -> Optional[str]:
    """
    Queries Google Gemini API using GEMINI_API_KEY from environment.
    Tries Google GenAI SDK first, then falls back to direct REST endpoints
    across models: gemini-2.5-flash, gemini-2.0-flash, gemini-1.5-flash.
    """
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        return None

    # Priority model sequence
    models = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-1.5-flash"]

    # 1. Try google-genai official client
    try:
        from google import genai
        client = genai.Client(api_key=api_key)
        for model_name in models:
            try:
                config = {"temperature": temperature}
                if json_mode:
                    config["response_mime_type"] = "application/json"
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                    config=config
                )
                if response and response.text:
                    return response.text
            except Exception as model_err:
                err_str = str(model_err)
                print(f"[Gemini SDK {model_name} Error]: {model_err}")
                if "401" in err_str or "UNAUTHENTICATED" in err_str or "API_KEY_INVALID" in err_str:
                    print("[Gemini API Notice]: Invalid/unauthenticated key. Using fast deterministic transformation fallback.")
                    return None
                continue
    except Exception as sdk_err:
        print(f"[Gemini SDK Initialization Error]: {sdk_err}")

    # 2. REST API fallback
    headers = {"Content-Type": "application/json"}
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": temperature
        }
    }
    if json_mode:
        payload["generationConfig"]["response_mime_type"] = "application/json"

    for model_name in models:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
        try:
            res = requests.post(url, headers=headers, json=payload, timeout=20)
            if res.status_code == 200:
                res_data = res.json()
                candidates = res_data.get("candidates", [])
                if candidates:
                    text = candidates[0].get("content", {}).get("parts", [{}])[0].get("text", "")
                    if text:
                        return text
            elif res.status_code == 401 or "UNAUTHENTICATED" in res.text or "API_KEY_INVALID" in res.text:
                print(f"[Gemini REST HTTP 401]: Invalid key. Using fast deterministic transformation fallback.")
                return None
            else:
                print(f"[Gemini REST {model_name} HTTP {res.status_code}]: {res.text[:150]}")
        except Exception as rest_err:
            print(f"[Gemini REST {model_name} Error]: {rest_err}")
            continue

    return None


# ==============================================================================
# Format 1: SUMMARY
# ==============================================================================

def generate_summary(text: str, audience: str = "General", tone: str = "Professional", length: str = "Medium") -> Dict[str, Any]:
    """Generates a concise, structured executive summary."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's expert content summarizer.
Produce a concise, structured executive briefing strictly grounded in the provided source text.
Audience: {audience} | Tone: {tone} | Target Length: {length}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "title": "Descriptive Executive Summary Title",
  "summary": "2-3 dense, articulate paragraphs providing a concise structured summary.",
  "highlights": [
    "Key highlight point 1",
    "Key highlight point 2",
    "Key highlight point 3",
    "Key highlight point 4"
  ],
  "key_takeaways": [
    "Strategic takeaway 1",
    "Strategic takeaway 2",
    "Strategic takeaway 3"
  ]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.2)
    if raw:
        try:
            cleaned = clean_json_response(raw)
            data = json.loads(cleaned)
            return {
                "format": "summary",
                "content": data,
                "status": "success",
                "message": "Structured summary generated successfully by Gemini."
            }
        except Exception as e:
            print(f"[Summary JSON Parse Error]: {e}")

    # Deterministic Fallback
    summary_sentences = meta["sentences"][:4] if len(meta["sentences"]) >= 4 else meta["sentences"]
    return {
        "format": "summary",
        "content": {
            "title": f"Executive Briefing: {meta['main_topic']} Overview",
            "summary": " ".join(summary_sentences) if summary_sentences else meta["first_sentence"],
            "highlights": meta["sentences"][1:5] if len(meta["sentences"]) > 4 else [meta["first_sentence"]],
            "key_takeaways": [
                f"Core focus on {meta['main_topic']} with vital data points preserved from source.",
                "Review detailed sections below to coordinate execution steps.",
                "Action items and operational guidelines highlighted across documentation."
            ]
        },
        "status": "fallback",
        "message": "Generated structured summary using source text analysis."
    }


# ==============================================================================
# Format 2: BLOG
# ==============================================================================

def generate_blog(text: str, audience: str = "General", tone: str = "Engaging", length: str = "Medium") -> Dict[str, Any]:
    """Generates an engaging blog post with title, introduction, sections, and conclusion."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's expert blog author.
Write a publication-ready blog post based strictly on the source content.
Target Audience: {audience} | Tone: {tone} | Length: {length}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "title": "Catchy, High-CTR Blog Post Title",
  "read_time": "4 min read",
  "tags": ["Topic1", "Topic2", "Topic3"],
  "introduction": "Compelling hook and opening paragraphs introducing the context and why it matters.",
  "sections": [
    {{
      "heading": "Section 1 Subheading",
      "content": "Detailed explanatory content covering this aspect of the source material."
    }},
    {{
      "heading": "Section 2 Subheading",
      "content": "Detailed analysis with facts, numbers, and practical context."
    }},
    {{
      "heading": "Section 3 Subheading",
      "content": "Actionable insights, implications, and key perspectives."
    }}
  ],
  "conclusion": "Thoughtful concluding summary with final thoughts or forward-looking call-to-action."
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.5)
    if raw:
        try:
            cleaned = clean_json_response(raw)
            data = json.loads(cleaned)
            return {
                "format": "blog",
                "content": data,
                "status": "success",
                "message": "Blog post generated successfully by Gemini."
            }
        except Exception as e:
            print(f"[Blog JSON Parse Error]: {e}")

    # Deterministic Fallback
    s = meta["sentences"]
    intro = s[0] + " " + (s[1] if len(s) > 1 else "")
    sec1 = s[2] if len(s) > 2 else "Examining the fundamental components of this domain."
    sec2 = s[3] if len(s) > 3 else "Assessing implications and data-driven insights from the source."
    sec3 = s[4] if len(s) > 4 else "Applying recommendations to optimize operational workflows."
    conc = s[5] if len(s) > 5 else f"In summary, understanding {meta['main_topic']} empowers informed strategic decisions."

    return {
        "format": "blog",
        "content": {
            "title": f"Exploring {meta['main_topic']}: Key Insights & Practical Perspectives",
            "read_time": "3 min read",
            "tags": [meta["main_topic"], "Insights", "Analysis"],
            "introduction": intro,
            "sections": [
                {"heading": f"Understanding the Foundations of {meta['main_topic']}", "content": sec1},
                {"heading": "Key Analysis and Practical Evidence", "content": sec2},
                {"heading": "Strategic Next Steps & Practical Guidance", "content": sec3}
            ],
            "conclusion": conc
        },
        "status": "fallback",
        "message": "Generated blog post using source text analysis."
    }


# ==============================================================================
# Format 3: MCQ
# ==============================================================================

def generate_mcq(text: str, audience: str = "Student", tone: str = "Educational", length: str = "Medium") -> Dict[str, Any]:
    """Generates an assessment quiz with questions, 4 options, correct answer, and explanation."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's educational assessment creator.
Create a high-quality 5-question multiple choice assessment based strictly on the factual details in the source text.
Audience: {audience} | Tone: {tone}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RULES:
1. Create 5 clear, objective questions testing genuine comprehension.
2. Each question MUST have exactly 4 options prefixed with 'A) ', 'B) ', 'C) ', 'D) '.
3. Only ONE option should be unambiguously correct.
4. 'correct_answer' must be the exact single letter 'A', 'B', 'C', or 'D'.
5. 'explanation' must explain WHY the correct answer is right and clarify any common misconceptions.

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "quiz_title": "Assessment: {meta['main_topic']} Comprehension Check",
  "questions": [
    {{
      "id": 1,
      "question": "Question text here?",
      "options": [
        "A) Option 1",
        "B) Option 2",
        "C) Option 3",
        "D) Option 4"
      ],
      "correct_answer": "A",
      "explanation": "Detailed explanation clarifying why Option A is supported directly by the source."
    }}
  ]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.2)
    if raw:
        try:
            cleaned = clean_json_response(raw)
            data = json.loads(cleaned)
            return {
                "format": "mcq",
                "content": data,
                "status": "success",
                "message": "MCQ quiz generated successfully by Gemini."
            }
        except Exception as e:
            print(f"[MCQ JSON Parse Error]: {e}")

    # Deterministic Fallback
    fallback_questions = []
    sample_s = meta["sentences"][:4] if meta["sentences"] else ["ContentForge AI enables rapid multi-format transformation."]
    for idx, sentence in enumerate(sample_s):
        fallback_questions.append({
            "id": idx + 1,
            "question": f"Based on the source document, which statement regarding {meta['main_topic']} is directly supported?",
            "options": [
                f"A) {sentence}",
                f"B) The source explicitly contradicts {meta['main_topic']} principles.",
                "C) No metrics or factual claims were stated in the documentation.",
                "D) The primary topic is unrelated to the discussed parameters."
            ],
            "correct_answer": "A",
            "explanation": f"Option A directly reflects the source statement: \"{sentence[:100]}...\""
        })

    return {
        "format": "mcq",
        "content": {
            "quiz_title": f"Knowledge Assessment: {meta['main_topic']}",
            "questions": fallback_questions
        },
        "status": "fallback",
        "message": "Generated knowledge assessment using source text analysis."
    }


# ==============================================================================
# Format 4: IMAGE
# ==============================================================================

def generate_image_asset(text: str, audience: str = "General", tone: str = "Creative", length: str = "Medium") -> Dict[str, Any]:
    """
    Generates an optimized image prompt using Gemini,
    and connects it to an image generation API (Pollinations Flux) for live visual generation.
    """
    meta = extract_quick_metadata(text)
    prompt = f"""You are an elite AI art director and prompt engineer.
Analyze the source text and craft an ultra-detailed, vivid image generation prompt that encapsulates the main subject, mood, and visual narrative of the source content.
Audience: {audience} | Tone: {tone}

SOURCE TEXT:
\"\"\"
{text[:8000]}
\"\"\"

RULES:
1. 'image_prompt': A rich 40-70 word prompt tailored for modern diffusion models (e.g. Flux, Midjourney). Include specific visual subject, atmosphere, dramatic cinematic lighting, composition, depth of field, 8k resolution, modern aesthetic. Do NOT include words like 'text' or 'watermark'.
2. 'style': The visual genre (e.g., 'Cinematic Digital Art', 'Photorealistic Editorial', 'Minimalist Isometric 3D', 'Architectural Visualization').
3. 'aspect_ratio': '16:9'
4. 'description': Brief 2-sentence rationale explaining the artistic symbolism and connection to the document.

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "image_prompt": "Cinematic visual prompt...",
  "style": "Cinematic Digital Art",
  "aspect_ratio": "16:9",
  "description": "Symbolic visual representation of the core themes..."
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.7)
    image_prompt = None
    style = "Cinematic Digital Art"
    aspect_ratio = "16:9"
    description = f"Artistic concept visualizing the core themes of {meta['main_topic']}."

    if raw:
        try:
            cleaned = clean_json_response(raw)
            data = json.loads(cleaned)
            image_prompt = data.get("image_prompt")
            style = data.get("style", style)
            aspect_ratio = data.get("aspect_ratio", aspect_ratio)
            description = data.get("description", description)
        except Exception as e:
            print(f"[Image JSON Parse Error]: {e}")

    if not image_prompt:
        image_prompt = f"Stunning cinematic digital illustration depicting {meta['main_topic']}, modern technological concepts, volumetric lighting, rich color palette, intricate detail, photorealistic 8k render, masterpiece composition, award winning visual"

    # Connect to live image generation model API (Pollinations Flux engine)
    encoded_prompt = urllib.parse.quote(image_prompt)
    image_url = f"https://image.pollinations.ai/prompt/{encoded_prompt}?width=1024&height=576&model=flux&nologo=true&seed={int(time.time()) % 10000}"

    return {
        "format": "image",
        "content": {
            "image_prompt": image_prompt,
            "style": style,
            "aspect_ratio": aspect_ratio,
            "image_url": image_url,
            "description": description,
            "download_url": image_url
        },
        "status": "success" if raw else "fallback",
        "message": "Optimized AI visual prompt crafted and image generation link created."
    }


# ==============================================================================
# Format 5: DOC (.docx via python-docx)
# ==============================================================================

def create_docx_document(structured_doc: Dict[str, Any]) -> bytes:
    """Builds a polished, professional Microsoft Word (.docx) document using python-docx."""
    doc = docx.Document()

    # Set standard 1-inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Document Header Title
    title_text = structured_doc.get("document_title", "Executive Report")
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(title_text)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900

    # Subtitle / Metadata
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run(f"Generated by ContentForge AI  •  {datetime.now().strftime('%B %d, %Y')}")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)  # Slate 500

    # Horizontal Divider Line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(14)
    div_run = div_p.add_run("_________________________________________________________________________________")
    div_run.font.size = Pt(8)
    div_run.font.color.rgb = RGBColor(226, 232, 240)

    # Executive Overview Section
    exec_summary = structured_doc.get("executive_summary", "")
    if exec_summary:
        h1 = doc.add_heading(level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(6)
        r_h1 = h1.add_run("Executive Overview")
        r_h1.font.name = "Calibri"
        r_h1.font.size = Pt(15)
        r_h1.font.bold = True
        r_h1.font.color.rgb = RGBColor(79, 70, 229)  # Indigo 600

        p_exec = doc.add_paragraph(exec_summary)
        p_exec.paragraph_format.space_after = Pt(12)
        p_exec.paragraph_format.line_spacing = 1.15
        for r in p_exec.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(51, 65, 85)

    # Document Body Sections
    for sec in structured_doc.get("sections", []):
        sec_title = sec.get("heading", "Topic Overview")
        h2 = doc.add_heading(level=2)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(4)
        r_h2 = h2.add_run(sec_title)
        r_h2.font.name = "Calibri"
        r_h2.font.size = Pt(13)
        r_h2.font.bold = True
        r_h2.font.color.rgb = RGBColor(30, 41, 59)  # Slate 800

        # Paragraphs
        for para_text in sec.get("paragraphs", []):
            p = doc.add_paragraph(para_text)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(51, 65, 85)

        # Bullet points
        for bullet_text in sec.get("bullet_points", []):
            bp = doc.add_paragraph(bullet_text, style="List Bullet")
            bp.paragraph_format.space_after = Pt(4)
            for r in bp.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(51, 65, 85)

    # Recommendations / Action Items
    recommendations = structured_doc.get("recommendations", [])
    if recommendations:
        h_rec = doc.add_heading(level=1)
        h_rec.paragraph_format.space_before = Pt(16)
        h_rec.paragraph_format.space_after = Pt(6)
        r_rec = h_rec.add_run("Strategic Recommendations & Action Items")
        r_rec.font.name = "Calibri"
        r_rec.font.size = Pt(15)
        r_rec.font.bold = True
        r_rec.font.color.rgb = RGBColor(79, 70, 229)

        for item in recommendations:
            bp = doc.add_paragraph(item, style="List Bullet")
            bp.paragraph_format.space_after = Pt(4)
            for r in bp.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_doc_asset(text: str, audience: str = "Executive", tone: str = "Professional", length: str = "Medium") -> Dict[str, Any]:
    """Generates formatted document content using Gemini and exports a downloadable .docx using python-docx."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's executive document drafting engine.
Create a comprehensive, formal business/technical briefing document based on the source text.
Audience: {audience} | Tone: {tone} | Length: {length}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "document_title": "Formal Document Title",
  "subtitle": "Comprehensive Briefing & Strategic Overview",
  "executive_summary": "Thorough 2-3 paragraph executive digest synthesizing core findings.",
  "sections": [
    {{
      "heading": "Section 1 Subject",
      "paragraphs": [
        "Paragraph 1 text detailing context.",
        "Paragraph 2 text detailing evidence."
      ],
      "bullet_points": [
        "Key specific data point or takeaway",
        "Key operational fact"
      ]
    }},
    {{
      "heading": "Section 2 Subject",
      "paragraphs": [
        "Paragraph 1 analytical breakdown."
      ],
      "bullet_points": [
        "Operational metric or observation"
      ]
    }}
  ],
  "recommendations": [
    "Actionable next step 1",
    "Actionable next step 2",
    "Actionable next step 3"
  ]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.2)
    doc_data = None
    if raw:
        try:
            cleaned = clean_json_response(raw)
            doc_data = json.loads(cleaned)
        except Exception as e:
            print(f"[DOC JSON Parse Error]: {e}")

    if not doc_data:
        s = meta["sentences"]
        doc_data = {
            "document_title": f"Strategic Analysis: {meta['main_topic']}",
            "subtitle": "Executive Briefing and Operational Synthesis",
            "executive_summary": " ".join(s[:3]) if len(s) >= 3 else meta["first_sentence"],
            "sections": [
                {
                    "heading": f"Context & Background on {meta['main_topic']}",
                    "paragraphs": [s[3] if len(s) > 3 else "Reviewing key source parameters."],
                    "bullet_points": meta["keywords"][:4]
                },
                {
                    "heading": "Analytical Observations & Findings",
                    "paragraphs": [s[4] if len(s) > 4 else "Preserving core factual records from source documentation."],
                    "bullet_points": s[5:7] if len(s) > 6 else ["Direct compliance with stated guidelines"]
                }
            ],
            "recommendations": [
                "Establish structured review protocol for outlined findings.",
                "Incorporate factual parameters into tactical implementation plan.",
                "Monitor milestone developments and align stakeholder resources."
            ]
        }

    # Build .docx binary
    docx_bytes = create_docx_document(doc_data)
    filename = f"ContentForge_{meta['main_topic'].replace(' ', '_')}.docx"
    file_id = store_file(docx_bytes, filename, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    file_base64 = base64.b64encode(docx_bytes).decode("utf-8")

    return {
        "format": "doc",
        "content": {
            "document_title": doc_data.get("document_title"),
            "subtitle": doc_data.get("subtitle"),
            "executive_summary": doc_data.get("executive_summary"),
            "sections": doc_data.get("sections", []),
            "recommendations": doc_data.get("recommendations", []),
            "filename": filename,
            "file_id": file_id,
            "download_url": f"/api/download/{file_id}",
            "file_base64": file_base64,
            "file_size_bytes": len(docx_bytes)
        },
        "status": "success" if raw else "fallback",
        "message": "Word document (.docx) crafted and prepared for download."
    }


# ==============================================================================
# Format 6: PRESENTATION (.pptx via python-pptx)
# ==============================================================================

def create_pptx_presentation(slides_data: List[Dict[str, Any]], deck_title: str) -> bytes:
    """Builds a high-impact, modern 16:9 widescreen presentation using python-pptx."""
    prs = pptx.Presentation()
    # Configure 16:9 widescreen dimensions
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    c_primary = PptxRGBColor(79, 70, 229)    # Indigo
    c_dark = PptxRGBColor(15, 23, 42)        # Slate 900
    c_muted = PptxRGBColor(100, 116, 139)    # Slate 500

    for idx, slide_info in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        is_title_slide = (idx == 0)

        if is_title_slide:
            # Decorative top accent banner
            banner = slide.shapes.add_shape(
                1,  # MSO_SHAPE.RECTANGLE
                PptxInches(0), PptxInches(0), PptxInches(13.333), PptxInches(0.4)
            )
            banner.fill.solid()
            banner.fill.fore_color.rgb = c_primary
            banner.line.color.rgb = c_primary

            # Main Presentation Title
            tb = slide.shapes.add_textbox(PptxInches(1.2), PptxInches(2.2), PptxInches(10.9), PptxInches(2.5))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_info.get("title", deck_title)
            p.font.name = "Calibri"
            p.font.size = PptxPt(44)
            p.font.bold = True
            p.font.color.rgb = c_dark

            # Subtitle
            sub_text = slide_info.get("subtitle", "Executive Slide Deck  •  ContentForge AI")
            p_sub = tf.add_paragraph()
            p_sub.text = sub_text
            p_sub.font.name = "Calibri"
            p_sub.font.size = PptxPt(20)
            p_sub.font.color.rgb = c_muted
            p_sub.space_before = PptxPt(16)

            # Footer Metadata
            tb_foot = slide.shapes.add_textbox(PptxInches(1.2), PptxInches(6.2), PptxInches(10.9), PptxInches(0.8))
            p_foot = tb_foot.text_frame.paragraphs[0]
            p_foot.text = f"Prepared with ContentForge AI  |  {datetime.now().strftime('%B %Y')}"
            p_foot.font.size = PptxPt(11)
            p_foot.font.color.rgb = c_muted

        else:
            # Top Header Bar
            header_bar = slide.shapes.add_shape(
                1, PptxInches(0), PptxInches(0), PptxInches(13.333), PptxInches(0.2)
            )
            header_bar.fill.solid()
            header_bar.fill.fore_color.rgb = c_primary
            header_bar.line.color.rgb = c_primary

            # Slide Number Badge & Category
            tb_cat = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(0.55), PptxInches(11.3), PptxInches(0.4))
            p_cat = tb_cat.text_frame.paragraphs[0]
            p_cat.text = f"CONTENTFORGE BRIEFING  •  SLIDE {idx + 1} OF {len(slides_data)}"
            p_cat.font.size = PptxPt(10)
            p_cat.font.bold = True
            p_cat.font.color.rgb = c_primary

            # Slide Title
            tb_title = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(0.9), PptxInches(11.3), PptxInches(1.0))
            p_title = tb_title.text_frame.paragraphs[0]
            p_title.text = slide_info.get("title", f"Slide {idx + 1}")
            p_title.font.name = "Calibri"
            p_title.font.size = PptxPt(30)
            p_title.font.bold = True
            p_title.font.color.rgb = c_dark

            # Content Box Card
            bullets = slide_info.get("bullets", [])
            tb_content = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(2.2), PptxInches(11.3), PptxInches(4.5))
            tf_content = tb_content.text_frame
            tf_content.word_wrap = True

            for b_idx, bullet in enumerate(bullets):
                p_b = tf_content.paragraphs[0] if b_idx == 0 else tf_content.add_paragraph()
                p_b.text = f"•  {bullet}"
                p_b.font.name = "Calibri"
                p_b.font.size = PptxPt(18)
                p_b.font.color.rgb = c_dark
                p_b.space_before = PptxPt(12)
                p_b.space_after = PptxPt(4)

        # Attach speaker notes if present
        notes_text = slide_info.get("notes", "")
        if notes_text:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = notes_text

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def generate_presentation_asset(text: str, audience: str = "Executive", tone: str = "Professional", length: str = "Medium") -> Dict[str, Any]:
    """Generates slide-wise content using Gemini and exports a downloadable .pptx presentation using python-pptx."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's executive presentation designer.
Generate a structured 5-6 slide presentation deck based strictly on the source text.
Audience: {audience} | Tone: {tone} | Length: {length}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RULES:
1. Slide 1 must be the Title Slide (include 'title' and 'subtitle').
2. Slide 2 should be the Executive Agenda / Key Themes.
3. Slide 3-5 should cover core topic deep dives with 3-4 impactful bullet points each.
4. Final Slide should be Strategic Conclusions & Actionable Next Steps.
5. Provide concise 'notes' for speaker talking points on each slide.

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "presentation_title": "Presentation Main Title",
  "slides": [
    {{
      "slide_number": 1,
      "title": "Main Presentation Title",
      "subtitle": "Strategic Executive Briefing",
      "bullets": [],
      "notes": "Introductory remarks welcoming stakeholders."
    }},
    {{
      "slide_number": 2,
      "title": "Agenda & Core Themes",
      "subtitle": "",
      "bullets": [
        "Context and background",
        "Primary analytical findings",
        "Strategic next steps and milestones"
      ],
      "notes": "Outline the presentation roadmap."
    }},
    {{
      "slide_number": 3,
      "title": "Key Topic Analysis",
      "subtitle": "",
      "bullets": [
        "Detail 1 with metrics or facts",
        "Detail 2 with operational insight",
        "Detail 3 with risk or opportunity factor"
      ],
      "notes": "Discuss detailed findings from the source."
    }},
    {{
      "slide_number": 4,
      "title": "Strategic Next Steps & Wrap-Up",
      "subtitle": "",
      "bullets": [
        "Immediate priority 1",
        "Operational milestone 2",
        "Long term outcome"
      ],
      "notes": "Closing synthesis and call for questions."
    }}
  ]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.3)
    pres_data = None
    if raw:
        try:
            cleaned = clean_json_response(raw)
            pres_data = json.loads(cleaned)
        except Exception as e:
            print(f"[Presentation JSON Parse Error]: {e}")

    if not pres_data:
        s = meta["sentences"]
        pres_data = {
            "presentation_title": f"{meta['main_topic']}: Strategic Overview",
            "slides": [
                {
                    "slide_number": 1,
                    "title": f"{meta['main_topic']} Overview",
                    "subtitle": "Executive Transformation Briefing  •  ContentForge AI",
                    "bullets": [],
                    "notes": "Welcome team and review objective."
                },
                {
                    "slide_number": 2,
                    "title": "Core Themes & Context",
                    "subtitle": "",
                    "bullets": [
                        meta["first_sentence"],
                        f"Focused analysis on {meta['main_topic']} domain fundamentals.",
                        "Direct extraction of key facts, dates, and operational records."
                    ],
                    "notes": "Establish factual baseline."
                },
                {
                    "slide_number": 3,
                    "title": "Key Observations & Findings",
                    "subtitle": "",
                    "bullets": s[1:4] if len(s) > 3 else [f"Verified evidence for {meta['main_topic']}"],
                    "notes": "Walk through primary supporting evidence."
                },
                {
                    "slide_number": 4,
                    "title": "Strategic Conclusions & Next Steps",
                    "subtitle": "",
                    "bullets": [
                        "Synthesize cross-functional takeaways across stakeholders.",
                        "Execute on action items outlined in documentation.",
                        "Deploy transformed assets across relevant communication channels."
                    ],
                    "notes": "Conclude and open floor for strategic discussion."
                }
            ]
        }

    slides = pres_data.get("slides", [])
    deck_title = pres_data.get("presentation_title", f"{meta['main_topic']} Deck")

    # Build .pptx binary
    pptx_bytes = create_pptx_presentation(slides, deck_title)
    filename = f"ContentForge_{meta['main_topic'].replace(' ', '_')}.pptx"
    file_id = store_file(pptx_bytes, filename, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    file_base64 = base64.b64encode(pptx_bytes).decode("utf-8")

    return {
        "format": "presentation",
        "content": {
            "presentation_title": deck_title,
            "slide_count": len(slides),
            "slides": slides,
            "filename": filename,
            "file_id": file_id,
            "download_url": f"/api/download/{file_id}",
            "file_base64": file_base64,
            "file_size_bytes": len(pptx_bytes)
        },
        "status": "success" if raw else "fallback",
        "message": "Presentation (.pptx) deck crafted and prepared for download."
    }


# ==============================================================================
# Format 7: FLASHCARDS
# ==============================================================================

def generate_flashcards(text: str, audience: str = "Student", tone: str = "Educational", length: str = "Medium") -> Dict[str, Any]:
    """Generates study flashcards with front (question/term) and back (answer/definition)."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's flashcard creator.
Create 8 high-quality study flashcards based strictly on the source text.
Audience: {audience} | Tone: {tone}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "deck_title": "Flashcard Deck: {meta['main_topic']}",
  "card_count": 8,
  "cards": [
    {{
      "id": 1,
      "front": "Key term or question from the source",
      "back": "Clear, concise definition or answer based on the source",
      "category": "Concept"
    }}
  ]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.2)
    if raw:
        try:
            data = json.loads(clean_json_response(raw))
            return {"format": "flashcards", "content": data, "status": "success", "message": "Flashcards generated by Gemini."}
        except Exception as e:
            print(f"[Flashcards JSON Error]: {e}")

    s = meta["sentences"]
    cards = []
    for i, sentence in enumerate(s[:8]):
        cards.append({"id": i+1, "front": f"What does the source say about point {i+1}?", "back": sentence, "category": "Fact"})
    return {
        "format": "flashcards",
        "content": {"deck_title": f"Study Deck: {meta['main_topic']}", "card_count": len(cards), "cards": cards},
        "status": "fallback",
        "message": "Flashcards generated from source text."
    }


# ==============================================================================
# Format 8: VIDEO SCRIPT
# ==============================================================================

def generate_video(text: str, audience: str = "General", tone: str = "Engaging", length: str = "Medium") -> Dict[str, Any]:
    """Generates a structured video script with scenes, narration, and visuals."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's video script writer.
Write a compelling, structured video script based on the source content.
Audience: {audience} | Tone: {tone} | Length: {length}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "video_title": "Engaging Video Title",
  "duration_estimate": "3-5 minutes",
  "style": "Explainer / Documentary / Promotional",
  "hook": "Opening 10-second attention-grabbing hook line",
  "scenes": [
    {{
      "scene_number": 1,
      "title": "Scene Title",
      "narration": "Full narration text for this scene.",
      "visual_direction": "Description of what appears on screen.",
      "duration": "30 seconds"
    }}
  ],
  "call_to_action": "Closing call-to-action line"
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.4)
    if raw:
        try:
            data = json.loads(clean_json_response(raw))
            return {"format": "video", "content": data, "status": "success", "message": "Video script generated by Gemini."}
        except Exception as e:
            print(f"[Video JSON Error]: {e}")

    s = meta["sentences"]
    scenes = [
        {"scene_number": 1, "title": "Introduction", "narration": s[0] if s else meta["first_sentence"], "visual_direction": f"Opening shot related to {meta['main_topic']}", "duration": "30 seconds"},
        {"scene_number": 2, "title": "Core Content", "narration": " ".join(s[1:3]) if len(s) > 2 else "Exploring the key concepts.", "visual_direction": "Animated infographic or relevant footage", "duration": "60 seconds"},
        {"scene_number": 3, "title": "Key Takeaways", "narration": " ".join(s[3:5]) if len(s) > 4 else "Summarizing the main points.", "visual_direction": "Bullet points appearing on screen", "duration": "45 seconds"},
        {"scene_number": 4, "title": "Conclusion", "narration": f"That's everything you need to know about {meta['main_topic']}.", "visual_direction": "Outro with logo and CTA", "duration": "20 seconds"},
    ]
    return {
        "format": "video",
        "content": {"video_title": f"{meta['main_topic']}: A Complete Overview", "duration_estimate": "3-4 minutes", "style": "Explainer", "hook": meta["first_sentence"], "scenes": scenes, "call_to_action": f"Learn more about {meta['main_topic']} today!"},
        "status": "fallback",
        "message": "Video script generated from source text."
    }


# ==============================================================================
# Format 9: SOCIAL MEDIA POST
# ==============================================================================

def generate_social_media(text: str, audience: str = "General", tone: str = "Engaging", length: str = "Short") -> Dict[str, Any]:
    """Generates platform-specific social media posts."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's social media strategist.
Create platform-optimized social media posts based on the source content.
Audience: {audience} | Tone: {tone}

SOURCE TEXT:
\"\"\"
{text[:8000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "topic": "Core topic of the posts",
  "platforms": {{
    "twitter": {{
      "post": "Tweet text under 280 characters with relevant hashtags",
      "hashtags": ["#tag1", "#tag2", "#tag3"]
    }},
    "linkedin": {{
      "post": "Professional LinkedIn post 150-300 words with insights and a question to drive engagement",
      "hashtags": ["#Professional", "#Insights"]
    }},
    "instagram": {{
      "caption": "Engaging Instagram caption with emojis and storytelling",
      "hashtags": ["#tag1", "#tag2", "#tag3", "#tag4", "#tag5"]
    }},
    "facebook": {{
      "post": "Conversational Facebook post that encourages sharing and comments",
      "hashtags": ["#tag1", "#tag2"]
    }}
  }}
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.6)
    if raw:
        try:
            data = json.loads(clean_json_response(raw))
            return {"format": "social_media", "content": data, "status": "success", "message": "Social media posts generated by Gemini."}
        except Exception as e:
            print(f"[Social Media JSON Error]: {e}")

    kw = meta["keywords"][:3]
    hashtags = [f"#{k.replace(' ', '')}" for k in kw]
    return {
        "format": "social_media",
        "content": {
            "topic": meta["main_topic"],
            "platforms": {
                "twitter": {"post": f"{meta['first_sentence'][:200]} {' '.join(hashtags)}", "hashtags": hashtags},
                "linkedin": {"post": f"Key insights on {meta['main_topic']}:\n\n{meta['first_sentence']}\n\nWhat are your thoughts?", "hashtags": hashtags},
                "instagram": {"caption": f"✨ {meta['first_sentence']} \n\n#ContentForge {' '.join(hashtags)}", "hashtags": hashtags + ["#ContentForge", "#AI"]},
                "facebook": {"post": f"Did you know? {meta['first_sentence']} Share this with someone who needs to see it!", "hashtags": hashtags}
            }
        },
        "status": "fallback",
        "message": "Social media posts generated from source text."
    }


# ==============================================================================
# Format 10: ADVERTISEMENT
# ==============================================================================

def generate_advertisement(text: str, audience: str = "General", tone: str = "Persuasive", length: str = "Short") -> Dict[str, Any]:
    """Generates multi-format advertisement copy."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's advertising copywriter.
Create compelling advertisement copy based on the source content.
Target Audience: {audience} | Tone: {tone}

SOURCE TEXT:
\"\"\"
{text[:8000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "product_service": "What is being advertised",
  "tagline": "Memorable 5-10 word tagline",
  "headline": "Attention-grabbing headline",
  "subheadline": "Supporting subheadline",
  "body_copy": "2-3 sentence persuasive body copy",
  "call_to_action": "Strong CTA button text",
  "google_ad": {{
    "headline_1": "Google Ad Headline 1 (max 30 chars)",
    "headline_2": "Google Ad Headline 2 (max 30 chars)",
    "description": "Google Ad description (max 90 chars)"
  }},
  "banner_ad": {{
    "headline": "Short banner headline",
    "body": "One line banner body",
    "cta": "CTA text"
  }},
  "key_benefits": ["Benefit 1", "Benefit 2", "Benefit 3"]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.6)
    if raw:
        try:
            data = json.loads(clean_json_response(raw))
            return {"format": "advertisement", "content": data, "status": "success", "message": "Advertisement copy generated by Gemini."}
        except Exception as e:
            print(f"[Advertisement JSON Error]: {e}")

    return {
        "format": "advertisement",
        "content": {
            "product_service": meta["main_topic"],
            "tagline": f"Discover the power of {meta['main_topic']}",
            "headline": f"Transform Your Understanding of {meta['main_topic']}",
            "subheadline": "Everything you need to know, simplified.",
            "body_copy": meta["first_sentence"],
            "call_to_action": "Learn More Today",
            "google_ad": {"headline_1": meta["main_topic"][:30], "headline_2": "Get Started Now", "description": meta["first_sentence"][:90]},
            "banner_ad": {"headline": meta["main_topic"], "body": "Explore key insights", "cta": "Discover Now"},
            "key_benefits": [f"Deep insights on {meta['main_topic']}", "Clear and actionable content", "Powered by AI"]
        },
        "status": "fallback",
        "message": "Advertisement copy generated from source text."
    }


# ==============================================================================
# Format 11: PROMOTIONAL CONTENT
# ==============================================================================

def generate_promotional(text: str, audience: str = "General", tone: str = "Enthusiastic", length: str = "Medium") -> Dict[str, Any]:
    """Generates promotional content including email, press release, and campaign copy."""
    meta = extract_quick_metadata(text)
    prompt = f"""You are ContentForge AI's promotional content specialist.
Create comprehensive promotional content based on the source material.
Audience: {audience} | Tone: {tone} | Length: {length}

SOURCE TEXT:
\"\"\"
{text[:10000]}
\"\"\"

RETURN JSON WITH THIS EXACT STRUCTURE:
{{
  "campaign_name": "Campaign or promotion name",
  "email_campaign": {{
    "subject_line": "Compelling email subject line",
    "preview_text": "Email preview text (50 chars)",
    "greeting": "Dear [Name],",
    "body": "2-3 paragraph promotional email body",
    "cta_text": "CTA button text",
    "sign_off": "Best regards, The Team"
  }},
  "press_release": {{
    "headline": "Press release headline",
    "dateline": "City, Date",
    "lead_paragraph": "First paragraph answering who, what, when, where, why",
    "body_paragraph": "Supporting details and context",
    "quote": "A relevant quote from a spokesperson",
    "boilerplate": "About ContentForge AI: AI-powered content transformation platform."
  }},
  "key_messages": ["Message 1", "Message 2", "Message 3"]
}}
"""
    raw = call_gemini_api(prompt, json_mode=True, temperature=0.5)
    if raw:
        try:
            data = json.loads(clean_json_response(raw))
            return {"format": "promotional", "content": data, "status": "success", "message": "Promotional content generated by Gemini."}
        except Exception as e:
            print(f"[Promotional JSON Error]: {e}")

    return {
        "format": "promotional",
        "content": {
            "campaign_name": f"{meta['main_topic']} Awareness Campaign",
            "email_campaign": {
                "subject_line": f"Exciting news about {meta['main_topic']}!",
                "preview_text": f"Key insights on {meta['main_topic']}",
                "greeting": "Dear Valued Reader,",
                "body": f"We are excited to share key insights about {meta['main_topic']}.\n\n{meta['first_sentence']}\n\nWe believe this information will be valuable to you.",
                "cta_text": "Read More",
                "sign_off": "Best regards,\nThe ContentForge AI Team"
            },
            "press_release": {
                "headline": f"New Insights on {meta['main_topic']} Released",
                "dateline": f"ContentForge AI, {datetime.now().strftime('%B %d, %Y')}",
                "lead_paragraph": meta["first_sentence"],
                "body_paragraph": " ".join(meta["sentences"][1:3]) if len(meta["sentences"]) > 2 else "Further details available upon request.",
                "quote": f"\"This content represents a significant development in {meta['main_topic']}.\"",
                "boilerplate": "About ContentForge AI: AI-powered multi-format content transformation platform."
            },
            "key_messages": [f"{meta['main_topic']} is a key area of focus", "AI-powered insights drive better decisions", "Transform content into actionable knowledge"]
        },
        "status": "fallback",
        "message": "Promotional content generated from source text."
    }


# ==============================================================================
# Master Transformation Dispatcher
# ==============================================================================

FORMAT_DISPATCHER = {
    "summary": generate_summary,
    "blog": generate_blog,
    "mcq": generate_mcq,
    "image": generate_image_asset,
    "doc": generate_doc_asset,
    "presentation": generate_presentation_asset,
    "flashcards": generate_flashcards,
    "video": generate_video,
    "social_media": generate_social_media,
    "advertisement": generate_advertisement,
    "promotional": generate_promotional,
}

SUPPORTED_FORMATS = list(FORMAT_DISPATCHER.keys())


def transform_content(
    source_text: str,
    selected_format: str,
    audience: str = "General",
    tone: str = "Professional",
    language: str = "English",
    length: str = "Medium"
) -> Dict[str, Any]:
    """
    Main entry point for Phase 5 content transformation:
    Dispatches to format-specific generator and returns standardized payload.
    """
    fmt_key = selected_format.lower().strip()
    if fmt_key not in FORMAT_DISPATCHER:
        raise ValueError(
            f"Unsupported format '{selected_format}'. Supported formats are: {', '.join(SUPPORTED_FORMATS)}"
        )

    generator_fn = FORMAT_DISPATCHER[fmt_key]
    result = generator_fn(
        text=source_text,
        audience=audience,
        tone=tone,
        length=length
    )
    return result
